"""Dokument-Parsing (AP 1.3) – pragmatischer v1-Backend-Satz.

Liefert eine einheitliche Block-Struktur (Text + optionale Seite/Überschrift),
auf der das Chunking arbeitet. Der Docling-Backend (Layout-/Tabellenerkennung,
ADR-6) und OCR für Scans docken später an derselben Schnittstelle an.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Block:
    text: str
    page: int | None = None
    heading_level: int | None = None  # 1–6, None = Fließtext


@dataclass
class ParsedDocument:
    title: str | None
    blocks: list[Block] = field(default_factory=list)


SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".html", ".htm", ".pdf", ".docx"}

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def parse_markdown(text: str) -> ParsedDocument:
    blocks: list[Block] = []
    paragraph: list[str] = []
    in_code = False

    def flush() -> None:
        if paragraph:
            blocks.append(Block(text="\n".join(paragraph).strip()))
            paragraph.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            paragraph.append(line)
            continue
        if in_code:
            paragraph.append(line)
            continue
        m = _HEADING_RE.match(line)
        if m:
            flush()
            blocks.append(Block(text=m.group(2).strip(), heading_level=len(m.group(1))))
        elif line.strip():
            paragraph.append(line)
        else:
            flush()
    flush()

    title = next((b.text for b in blocks if b.heading_level == 1), None)
    return ParsedDocument(title=title, blocks=[b for b in blocks if b.text])


def parse_text(text: str) -> ParsedDocument:
    blocks = [
        Block(text=p.strip()) for p in re.split(r"\n\s*\n", text) if p.strip()
    ]
    return ParsedDocument(title=None, blocks=blocks)


def parse_html(data: str) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()

    blocks: list[Block] = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre"]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name.startswith("h"):
            blocks.append(Block(text=text, heading_level=int(el.name[1])))
        else:
            blocks.append(Block(text=text))

    title_el = soup.find("title") or soup.find("h1")
    title = title_el.get_text(strip=True) if title_el else None
    return ParsedDocument(title=title, blocks=blocks)


def parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[Block] = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        for para in re.split(r"\n\s*\n", text):
            if para.strip():
                blocks.append(Block(text=para.strip(), page=page_no))
    meta_title = (reader.metadata or {}).get("/Title")
    title = str(meta_title).strip() if meta_title else None
    # Hinweis: Bild-PDFs (Scans) liefern hier keinen Text -> 0 Blöcke.
    # OCR-Backend (Tesseract/VLM, AP 1.3) wird an dieser Stelle ergänzt.
    return ParsedDocument(title=title or None, blocks=blocks)


_DOCX_HEADING_RE = re.compile(r"^Heading (\d)$|^Überschrift (\d)$")


def parse_docx(path: Path) -> ParsedDocument:
    import docx

    document = docx.Document(str(path))
    blocks: list[Block] = []
    title: str | None = document.core_properties.title or None
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        level: int | None = None
        style = para.style.name if para.style is not None else ""
        m = _DOCX_HEADING_RE.match(style or "")
        if m:
            level = int(m.group(1) or m.group(2))
            if title is None and level == 1:
                title = text
        blocks.append(Block(text=text, heading_level=level))
    return ParsedDocument(title=title, blocks=blocks)


def parse_file(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix in (".md", ".markdown"):
        return parse_markdown(path.read_text(encoding="utf-8", errors="replace"))
    if suffix == ".txt":
        return parse_text(path.read_text(encoding="utf-8", errors="replace"))
    if suffix in (".html", ".htm"):
        return parse_html(path.read_text(encoding="utf-8", errors="replace"))
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"Nicht unterstütztes Format: {suffix}")
